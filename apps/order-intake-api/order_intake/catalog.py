from __future__ import annotations

import csv
import io
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


COLUMN_ALIASES = {
    "sku": {"sku", "รหัส", "รหัสสินค้า", "itemcode", "productcode", "code"},
    "name": {"ชื่อสินค้า", "สินค้า", "ชื่อ", "product", "productname", "item", "itemname"},
    "aliases": {"ชื่อเรียก", "ชื่ออื่น", "alias", "aliases", "keyword", "keywords"},
    "uom": {"หน่วย", "หน่วยนับ", "uom", "unit"},
    "price": {"ราคา", "ราคาขาย", "price", "sellingprice", "rate"},
    "stock": {"คงเหลือ", "สต็อก", "stock", "stockqty", "qty", "quantity", "onhand"},
}


def normalize_header(value: Any) -> str:
    return re.sub(r"[^0-9a-zก-๙]", "", str(value or "").strip().lower())


def map_headers(values: list[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, value in enumerate(values):
        normalized = normalize_header(value)
        for field, aliases in COLUMN_ALIASES.items():
            if normalized in aliases and field not in mapped:
                mapped[field] = index
    return mapped


def _number(value: Any, default: str) -> str:
    if value is None or str(value).strip() == "":
        return default
    cleaned = re.sub(r"[^0-9.\-]", "", str(value).replace(",", ""))
    try:
        return format(Decimal(cleaned), "f")
    except (InvalidOperation, ValueError):
        return default


def _rows_to_products(rows: list[list[Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    best_row = -1
    best_mapping: dict[str, int] = {}
    for index, row in enumerate(rows[:20]):
        mapping = map_headers(row)
        if "name" in mapping and len(mapping) > len(best_mapping):
            best_row, best_mapping = index, mapping
    if best_row < 0:
        raise ValueError("could not find a product-name column in the first 20 rows")

    products: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[best_row + 1 :], start=best_row + 2):
        def cell(field: str, default=None):
            column = best_mapping.get(field)
            return row[column] if column is not None and column < len(row) else default

        name = str(cell("name") or "").strip()
        if not name:
            continue
        sku = str(cell("sku") or f"AUTO-{row_number:05d}").strip()
        aliases_value = str(cell("aliases") or "")
        aliases = [part.strip() for part in re.split(r"[,;|/]", aliases_value) if part.strip()]
        products.append(
            {
                "sku": sku,
                "name": name,
                "aliases": aliases,
                "uom": str(cell("uom") or "ชิ้น").strip(),
                "price": _number(cell("price"), "0"),
                "stock": _number(cell("stock"), "0"),
            }
        )
    return products, {
        "header_row": best_row + 1,
        "mapping": best_mapping,
        "product_count": len(products),
    }


def parse_catalog(data: bytes, filename: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = data.decode("utf-8-sig")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
        return _rows_to_products(rows)
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for XLSX imports") from exc
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        products, metadata = _rows_to_products(rows)
        metadata["sheet"] = sheet.title
        return products, metadata
    raise ValueError("catalog must be a .csv, .xlsx, or .xlsm file")
